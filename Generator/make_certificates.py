#!/usr/bin/env python3
"""
Certificate of Appreciation generator — 马来西亚原始佛教会

Usage:
    python3 make_certificates.py certificates.xlsx
    python3 make_certificates.py certificates.xlsx --only 3      # just row 3 of the sheet

Reads one row per certificate from the spreadsheet and writes, into ../Output/:
    <file name>.jpg   full resolution, 2480 x 3508 px, 300 dpi
    <file name>.pdf   A4 print-ready
    All certificates (<sheet name>).pdf   every certificate in one PDF, for printing in one go

Needs: numpy, Pillow, uharfbuzz, freetype-py, openpyxl
    pip install numpy pillow uharfbuzz freetype-py openpyxl
"""
import os, re, sys, datetime, argparse
import numpy as np
from PIL import Image
import uharfbuzz as hb
import freetype

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)                      # the OB-Cert folder
FONTS = os.path.join(ROOT, 'Fonts')
TEMPLATES = os.path.join(HERE, 'templates')
OUTPUT = os.path.join(ROOT, 'Output')

# ----------------------------------------------------------------------------------
# Settings taken from the original PSD files (300 dpi, so px = pt * 300 / 72)
# ----------------------------------------------------------------------------------
PT = 300 / 72
FONT_NAME = os.path.join(FONTS, 'Cinzel-Medium.ttf')                      # Cinzel Medium (current recipient font)
# Previous recipient font, kept for reference (use recipient_font = Brush Script.ttf in the sheet):
#   Brush Script MT, 30 pt single line; two lines approved at 20.6 pt / 26 pt
FONT_CN   = os.path.join(FONTS, 'BiMoChunQiu-ChaoQingKaiShu-2.ttf')      # AaBMCQCQKS (Aa筆墨春秋超清楷书)
FONT_EN   = os.path.join(FONTS, 'Apple Chancery.ttf')                    # Apple Chancery
FONT_DIGIT = os.path.join(FONTS, 'FreeSerif.ttf')                        # stands in for SIL Kai digits (layout A date)

BLACK = (1, 1, 1)
BROWN = (71, 46, 24)          # #472E18, all description text
DATE_BLACK = (0, 0, 0)

CENTRE_X = 1240.0             # visual centre line of the certificate
NAME_BASELINE = 1559.41       # baseline of the (last) recipient line, just above the gold line
NAME_TRACKING = 50
NAME_MAX_WIDTH = 1100         # the gold line under the name is 1149 px
DESC_MAX_WIDTH = 1420

CN_SIZE, CN_TRACKING = 17 * PT, 75
EN_SIZE, EN_TRACKING = 16 * PT, 75
EN_LEADING = 80               # Photoshop auto-leading, 120 % of 16 pt

CN_FIRST_LINE = '发心护持马来西亚原始佛教会'
EN_FIRST_LINE = 'In Sincere Appreciation of Your Invaluable'
EN_FIRST_LINE_CENTRE = 1219.0   # where the original artwork places this line

LAYOUTS = {
    # A: date on its own line under the signature (the "Durian Carnival" design)
    'A': dict(template='layout_A_date_below.png',
              cn_baselines=[1712.334, 1826.741],
              en_first_baseline=1962.088, en_max_lines=3),
    # B: date written into the Chinese and English text (the "Nestlé" design)
    'B': dict(template='layout_B_date_in_text.png',
              cn_baselines=[1712.334, 1815.741, 1919.741],
              en_first_baseline=2043.088, en_max_lines=3),
}

# ----------------------------------------------------------------------------------
# Text rendering (HarfBuzz shaping + FreeType, 4x supersampled), matches Photoshop
# ----------------------------------------------------------------------------------
SS = 4
_fonts = {}

def _font(path):
    if path not in _fonts:
        if not os.path.exists(path):
            sys.exit(f'Font file not found: {path}')
        _fonts[path] = (hb.Font(hb.Face(open(path, 'rb').read())), freetype.Face(path))
    return _fonts[path]

def _shape(text, path, size, tracking):
    hbf, _ = _font(path)
    sc = size / hbf.face.upem
    tr = tracking / 1000 * size
    buf = hb.Buffer(); buf.add_str(text); buf.guess_segment_properties()
    hb.shape(hbf, buf, {'kern': True, 'liga': True})
    glyphs, x = [], 0.0
    for info, pos in zip(buf.glyph_infos, buf.glyph_positions):
        glyphs.append((info.codepoint, x + pos.x_offset * sc, -pos.y_offset * sc))
        x += pos.x_advance * sc + tr
    return glyphs, x - tr

def text_width(text, path, size, tracking):
    return _shape(text, path, size, tracking)[1]

def render_line(W, H, text, path, size, tracking, x0, baseline):
    """Alpha mask (H, W) of one line, left edge of pen at x0."""
    glyphs, w = _shape(text, path, size, tracking)
    LX = max(int(x0 - size * 1.5), 0); RX = min(int(x0 + w + size * 1.5), W)
    TY = max(int(baseline - size * 1.5), 0); BY = min(int(baseline + size), H)
    lw, lh = RX - LX, BY - TY
    _, ft = _font(path)
    ft.set_char_size(int(size * SS * 64))
    big = np.zeros((lh * SS, lw * SS), np.float32)
    for gid, gx, gy in glyphs:
        px = (x0 + gx - LX) * SS; py = (baseline + gy - TY) * SS
        ix = int(np.floor(px)); fx = px - ix
        ft.set_transform(freetype.Matrix(0x10000, 0, 0, 0x10000), freetype.Vector(int(fx * 64), 0))
        ft.load_glyph(gid, freetype.FT_LOAD_RENDER | freetype.FT_LOAD_NO_HINTING)
        g = ft.glyph; bm = g.bitmap
        if bm.width == 0:
            continue
        arr = np.array(bm.buffer, np.float32).reshape(bm.rows, bm.pitch)[:, :bm.width] / 255
        X = ix + g.bitmap_left; Y = int(round(py)) - g.bitmap_top
        ys = slice(max(Y, 0), min(Y + bm.rows, lh * SS)); xs = slice(max(X, 0), min(X + bm.width, lw * SS))
        if ys.stop <= ys.start or xs.stop <= xs.start:
            continue
        big[ys, xs] = np.maximum(big[ys, xs], arr[ys.start - Y:ys.stop - Y, xs.start - X:xs.stop - X])
    out = np.zeros((H, W), np.float32)
    out[TY:BY, LX:RX] = big.reshape(lh, SS, lw, SS).mean(axis=(1, 3))
    return out

def centre_ink(mask, cx=CENTRE_X):
    """Shift a mask so its visible ink is centred on cx (whole-pixel shift)."""
    cols = np.where(mask.max(0) > 0.02)[0]
    if len(cols) == 0:
        return mask
    shift = int(round(cx - (cols[0] + cols[-1] + 1) / 2))
    return np.roll(mask, shift, axis=1)

def centred_line(W, H, text, path, size, tracking, baseline, max_width, warnings, label, cx=CENTRE_X):
    w = text_width(text, path, size, tracking)
    if w > max_width:
        new = size * max_width / w
        warnings.append(f'{label} "{text}" was too wide at {size / PT:.1f} pt, reduced to {new / PT:.1f} pt')
        size = new; w = max_width
    m = render_line(W, H, text, path, size, tracking, cx - w / 2, baseline)
    cols = np.where(m.max(0) > 0.02)[0]
    off = cx - (cols[0] + cols[-1] + 1) / 2          # centre the visible ink, to the sub-pixel
    return render_line(W, H, text, path, size, tracking, cx - w / 2 + off, baseline), size

def paint(img, mask, colour):
    return img * (1 - mask[..., None]) + np.array(colour, np.float32) * mask[..., None]

# ----------------------------------------------------------------------------------
# Dates
# ----------------------------------------------------------------------------------
MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
          'September', 'October', 'November', 'December']

def parse_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = str(v).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d %b %Y', '%d %B %Y', '%Y/%m/%d'):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f'Cannot read date "{s}" (use e.g. 2026-09-12 or 12 Sep 2026)')

def date_cn(d):
    return f'{d.year}年{d.month}月{d.day}日'

def date_en(d):
    n = d.day
    suf = 'th' if 11 <= n % 100 <= 13 else {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th')
    return f'{n}{suf} {MONTHS[d.month - 1]} {d.year}'

_cjk_date_glyphs = None

def render_date_A(W, H, d):
    """Layout A date line: digits in FreeSerif (≈ the original's digits), 年月日 taken from the original artwork."""
    global _cjk_date_glyphs
    if _cjk_date_glyphs is None:
        _cjk_date_glyphs = {
            '年': (np.array(Image.open(os.path.join(TEMPLATES, 'date_nian.png')), np.float32) / 255, 4.07),
            '月': (np.array(Image.open(os.path.join(TEMPLATES, 'date_yue.png')), np.float32) / 255, 7.55),
            '日': (np.array(Image.open(os.path.join(TEMPLATES, 'date_ri.png')), np.float32) / 255, 15.03),
        }
    text = date_cn(d)
    size = 54.16667; tr = 0.075 * size; top = 2664; baseline = 2707.19
    mask = np.zeros((H, W), np.float32)
    x = 1000.0
    for ch in text:
        if ch in _cjk_date_glyphs:
            g, bearing = _cjk_date_glyphs[ch]
            gx = x + bearing; ix = int(np.floor(gx)); fx = gx - ix
            gg = np.pad(g, ((0, 0), (0, 2)))
            gg = (1 - fx) * gg + fx * np.roll(gg, 1, 1)
            sl = mask[top:top + gg.shape[0], ix:ix + gg.shape[1]]
            np.maximum(sl, gg, out=sl)
            x += size + tr
        else:
            mask = np.maximum(mask, render_line(W, H, ch, FONT_DIGIT, 54.0, 0, x + 2.6, baseline))
            x += 0.5 * size + tr
    return centre_ink(mask)

# ----------------------------------------------------------------------------------
# One certificate
# ----------------------------------------------------------------------------------
def split_lines(s):
    return [p.strip() for p in str(s or '').split('|') if p.strip()]

def make_certificate(row):
    warnings = []
    layout = str(row.get('layout') or 'B').strip().upper()
    if layout not in LAYOUTS:
        raise ValueError(f'layout must be A or B, not "{layout}"')
    L = LAYOUTS[layout]
    base = Image.open(os.path.join(TEMPLATES, L['template'])).convert('RGB')
    W, H = base.size
    img = np.array(base).astype(np.float32)

    d = parse_date(row['date'])

    # --- recipient ---
    names = [n for n in [row.get('recipient_line_1'), row.get('recipient_line_2')] if n and str(n).strip()]
    if not names:
        raise ValueError('recipient_line_1 is empty')
    names = [str(n).strip() for n in names]
    # Cinzel defaults: 24 pt for one line; 20 pt over 22 pt for two lines (approved for DOB Design & Build)
    default_pts = [24] if len(names) == 1 else [20, 22]
    sizes = []
    for i, key in enumerate(['recipient_line_1_pt', 'recipient_line_2_pt'][:len(names)]):
        v = row.get(key)
        sizes.append(float(v) * PT if v not in (None, '') else default_pts[i] * PT)
    # last line sits on the template baseline; the line above sits 1.08 x (lower line's size) higher
    baselines = [NAME_BASELINE] if len(names) == 1 else [NAME_BASELINE - 1.08 * sizes[1], NAME_BASELINE]
    name_font = os.path.join(FONTS, str(row.get('recipient_font')).strip()) if row.get('recipient_font') else FONT_NAME
    for n, s, b in zip(names, sizes, baselines):
        m, _ = centred_line(W, H, n, name_font, s, NAME_TRACKING, b, NAME_MAX_WIDTH, warnings, 'Recipient')
        img = paint(img, m, BLACK)

    # --- Chinese description ---
    cn = [str(row.get('cn_line_1') or CN_FIRST_LINE).strip()]
    if layout == 'B':
        cn.append('于' + date_cn(d))
    cn += split_lines(row.get('event_cn'))
    if len(cn) > len(L['cn_baselines']):
        raise ValueError(f'too many Chinese lines for layout {layout}: {cn}')
    for t, b in zip(cn, L['cn_baselines']):
        m, _ = centred_line(W, H, t, FONT_CN, CN_SIZE, CN_TRACKING, b, DESC_MAX_WIDTH, warnings, 'Chinese line')
        img = paint(img, m, BROWN)

    # --- English description ---
    en = [str(row.get('en_line_1') or EN_FIRST_LINE).strip()] + split_lines(row.get('event_en'))
    if layout == 'B':
        # add "on <date>" to the last English line if it still fits, otherwise on its own line
        joined = en[-1] + ' on ' + date_en(d)
        if len(en) > 1 and text_width(joined, FONT_EN, EN_SIZE, EN_TRACKING) <= DESC_MAX_WIDTH:
            en[-1] = joined
        else:
            en.append('on ' + date_en(d))
    if len(en) > L['en_max_lines']:
        raise ValueError(f'too many English lines for layout {layout} (max {L["en_max_lines"]}): {en}')
    for i, t in enumerate(en):
        cx = EN_FIRST_LINE_CENTRE if (i == 0 and t == EN_FIRST_LINE) else CENTRE_X
        m, _ = centred_line(W, H, t, FONT_EN, EN_SIZE, EN_TRACKING, L['en_first_baseline'] + i * EN_LEADING,
                            DESC_MAX_WIDTH, warnings, 'English line', cx)
        img = paint(img, m, BROWN)

    # --- date under the signature (layout A only) ---
    if layout == 'A':
        img = paint(img, render_date_A(W, H, d), DATE_BLACK)

    return Image.fromarray(np.clip(img + 0.5, 0, 255).astype(np.uint8)), warnings

# ----------------------------------------------------------------------------------
# Spreadsheet
# ----------------------------------------------------------------------------------
COLUMNS = ['layout', 'recipient_line_1', 'recipient_line_2', 'recipient_line_1_pt', 'recipient_line_2_pt',
           'date', 'event_cn', 'event_en', 'cn_line_1', 'en_line_1', 'file_name', 'recipient_font']

def read_rows(path):
    if path.lower().endswith('.csv'):
        import csv
        with open(path, encoding='utf-8-sig') as f:
            rows = list(csv.DictReader(f))
        return [(i + 2, {k.strip(): v for k, v in r.items() if k}) for i, r in enumerate(rows)]
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)['Certificates']
    header = [str(c.value).strip() if c.value else '' for c in ws[1]]
    out = []
    for r in range(2, ws.max_row + 1):
        vals = {header[c]: ws.cell(r, c + 1).value for c in range(len(header)) if header[c]}
        if any(v not in (None, '') for v in vals.values()):
            out.append((r, vals))
    return out

def safe(s):
    return re.sub(r'[\\/:*?"<>|]+', '', s).strip()

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('sheet')
    ap.add_argument('--only', type=int, help='spreadsheet row number to generate')
    a = ap.parse_args()
    os.makedirs(OUTPUT, exist_ok=True)
    rows = read_rows(a.sheet)
    if a.only:
        rows = [r for r in rows if r[0] == a.only]
    pages, made = [], 0
    for rownum, row in rows:
        try:
            img, warns = make_certificate(row)
        except Exception as e:
            print(f'Row {rownum}: SKIPPED - {e}')
            continue
        name = row.get('file_name') or f"Cert of Appreciation-{row.get('recipient_line_1')}"
        name = safe(str(name))
        img.save(os.path.join(OUTPUT, name + '.jpg'), quality=99, subsampling=0, dpi=(300, 300))
        img.save(os.path.join(OUTPUT, name + '.pdf'), resolution=300)
        pages.append(img); made += 1
        print(f'Row {rownum}: {name}')
        for w in warns:
            print(f'    note: {w}')
    if len(pages) > 1:
        combined = os.path.join(OUTPUT, f'All certificates ({safe(os.path.splitext(os.path.basename(a.sheet))[0])}).pdf')
        pages[0].save(combined, resolution=300, save_all=True, append_images=pages[1:])
        print(f'Combined PDF: {os.path.basename(combined)}')
    print(f'Done: {made} certificate(s) in {OUTPUT}')

if __name__ == '__main__':
    main()
